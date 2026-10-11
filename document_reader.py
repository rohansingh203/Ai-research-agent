
from io import BytesIO, StringIO
from pathlib import Path
import csv

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".md", ".txt", ".csv", ".xlsx"
}


def read_document(filename: str, file_bytes: bytes) -> str:
    """Extract text while retaining source and PDF page information."""

    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        parts = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or "[No extractable text]"
            parts.append(
                f"[Source: {filename}, page {page_number}]\n{text}"
            )

        return "\n\n".join(parts)

    if suffix == ".docx":
        doc = Document(BytesIO(file_bytes))
        parts = [
            p.text for p in doc.paragraphs if p.text.strip()
        ]

        for table_number, table in enumerate(doc.tables, start=1):
            parts.append(f"Table {table_number}:")
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )

        return f"[Source: {filename}]\n" + "\n".join(parts)

    if suffix in {".md", ".txt"}:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        return f"[Source: {filename}]\n{text}"

    if suffix == ".csv":
        text = file_bytes.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(StringIO(text)))
        content = "\n".join(" | ".join(row) for row in rows)
        return f"[Source: {filename}]\n{content}"

    if suffix == ".xlsx":
        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
        parts = []

        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    "" if value is None else str(value)
                    for value in row
                ]

                if any(values):
                    parts.append(" | ".join(values))

        workbook.close()
        return f"[Source: {filename}]\n" + "\n".join(parts)

    raise ValueError(f"Unsupported file type: {suffix}")


def read_research_pack(uploaded_files) -> list[dict]:
    """Return source-labelled documents in upload order."""

    documents = []

    for index, uploaded_file in enumerate(uploaded_files, start=1):
        filename = uploaded_file.name
        suffix = Path(filename).suffix.lower()

        if suffix not in SUPPORTED_EXTENSIONS:
            continue

        content = read_document(filename, uploaded_file.getvalue())

        documents.append({
            "source_id": f"S{index}",
            "filename": filename,
            "content": content,
        })

    return documents